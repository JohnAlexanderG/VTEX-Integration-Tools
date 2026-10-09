#!/usr/bin/env python3
"""
Filtrar precios por listado de SKUs (CSV -> XLSX)

Toma un CSV con un listado de SKU IDs (p. ej. el `_skus_sin_precio.csv` de la
herramienta 76) y un xlsx de precios en formato VTEX, y exporta un xlsx con
solo los precios de esos SKUs, listo para importar en VTEX.

Uso:
    python3 filter_prices_by_sku.py <skus.csv> <precios.xlsx> <output_prefix> [--column "SKU ID"]

Ejemplo:
    python3 77_price_filter_by_sku_list/filter_prices_by_sku.py \\
        precios_erp_skus_sin_precio.csv dido-base-prices.xlsx output/didopet_skus_sin_precio

Entradas:
    - skus.csv:     CSV con columna SKU ID (configurable con --column)
    - precios.xlsx: SKU ID, Cost Price, Base Price, List Price, Error Code, Error Message
                    (se usa la primera hoja que tenga esas columnas)

Salidas:
    - {prefix}_precios.xlsx:         Precios de los SKUs del CSV (mismas 6 columnas)
    - {prefix}_no_encontrados.csv:   SKUs del CSV sin precio en el xlsx
    - {prefix}_REPORT.md:            Reporte con estadísticas
"""

import argparse
import csv
import os
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("❌ Error: falta openpyxl. Instalar con: pip install openpyxl")
    sys.exit(1)


OUTPUT_COLUMNS = ['SKU ID', 'Cost Price', 'Base Price', 'List Price', 'Error Code', 'Error Message']
REQUIRED_COLUMNS = ['SKU ID', 'Cost Price', 'Base Price']
HEADER_SCAN_ROWS = 10


def header_columns(row):
    return {str(name).strip(): idx for idx, name in enumerate(row or []) if name is not None and str(name).strip()}


def find_header(rows):
    """Devuelve (columnas, fila_encabezado) si alguna de las primeras filas tiene las columnas requeridas."""
    for header_row, row in enumerate(rows, start=1):
        candidate = header_columns(row)
        if all(c in candidate for c in REQUIRED_COLUMNS):
            return candidate, header_row
        if header_row >= HEADER_SCAN_ROWS:
            break
    return None, 0


def read_prices_xlsx(file_path):
    """
    Lee los precios de la primera hoja que contenga las columnas requeridas.

    Detecta la fila de encabezados dentro de las primeras HEADER_SCAN_ROWS para
    omitir filas de ayuda que VTEX agrega arriba.
    """
    if not os.path.exists(file_path):
        print(f"❌ Error: El archivo '{file_path}' no existe")
        sys.exit(1)

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        columns, header_row = find_header(rows)
        if columns is None:
            continue

        print(f"   ℹ️  Hoja '{ws.title}', encabezados en fila {header_row}")
        records = []
        for row in rows:
            record = {name: (row[idx] if idx < len(row) else None) for name, idx in columns.items()}
            if is_blank(record.get('SKU ID')):
                continue
            records.append(record)
        wb.close()
        return records, ws.title

    wb.close()
    print(f"❌ Error: ninguna hoja de '{file_path}' tiene las columnas: {', '.join(REQUIRED_COLUMNS)}")
    sys.exit(1)


def read_skus_csv(file_path, column):
    if not os.path.exists(file_path):
        print(f"❌ Error: El archivo '{file_path}' no existe")
        sys.exit(1)

    with open(file_path, encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f)
        fieldnames = [name.strip() for name in (reader.fieldnames or [])]
        reader.fieldnames = fieldnames
        if column not in fieldnames:
            print(f"❌ Error: '{file_path}' no tiene la columna '{column}'")
            print(f"   Columnas encontradas: {', '.join(fieldnames)}")
            sys.exit(1)
        return list(reader), fieldnames


def to_sku_id(value):
    """Normaliza SKU ID a int (el export de precios lo trae como float 1.0)."""
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def to_text(value):
    """Texto como en el export de VTEX: enteros sin decimales ('47300'), vacío si no hay valor."""
    if is_blank(value):
        return ''
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def is_blank(value):
    return value is None or str(value).strip() == ''


def write_prices_xlsx(file_path, records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Precios'
    ws.append(OUTPUT_COLUMNS)
    # Todo como texto, igual que el export de VTEX: con celdas numéricas la importación falla
    for record in records:
        row = [str(to_sku_id(record['SKU ID']))]
        row += [to_text(record.get(field)) for field in OUTPUT_COLUMNS[1:]]
        ws.append(row)
    wb.save(file_path)


def write_csv(file_path, fieldnames, rows):
    with open(file_path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)


def generate_report(report_file, stats, files, sheet_name):
    lines = [
        "# 💲 Filtro de precios por listado de SKUs",
        "",
        f"**Fecha:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "## 📊 Estadísticas",
        "",
        "| Métrica | Valor |",
        "|---|---|",
        f"| Filas de precios en xlsx (hoja `{sheet_name}`) | {stats['price_rows']} |",
        f"| SKU IDs duplicados en xlsx | {stats['price_duplicates']} |",
        f"| SKUs en CSV | {stats['csv_rows']} |",
        f"| SKU IDs inválidos en CSV | {stats['csv_invalid']} |",
        f"| ✅ Con precio (exportados) | {stats['found']} |",
        f"| ⚠️ Sin precio en xlsx | {stats['not_found']} |",
        "",
        "## 📁 Archivos generados",
        "",
    ]
    lines += [f"- `{os.path.basename(path)}`" for path in files]
    with open(report_file, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')


def main():
    parser = argparse.ArgumentParser(
        description='Exporta a xlsx solo los precios de los SKUs listados en un CSV.')
    parser.add_argument('skus_csv', help='CSV con el listado de SKUs')
    parser.add_argument('prices_xlsx', help='xlsx de precios (SKU ID, Cost Price, Base Price, List Price, ...)')
    parser.add_argument('output_prefix', help='Prefijo de los archivos de salida')
    parser.add_argument('--column', default='SKU ID', help='Columna del CSV con el SKU ID (default: "SKU ID")')
    args = parser.parse_args()

    print("💲 Filtro de precios por listado de SKUs")
    print(f"📥 Leyendo precios: {args.prices_xlsx}")
    price_records, sheet_name = read_prices_xlsx(args.prices_xlsx)

    prices_by_sku = {}
    price_duplicates = 0
    for record in price_records:
        sku_id = to_sku_id(record['SKU ID'])
        if sku_id is None:
            continue
        if sku_id in prices_by_sku:
            price_duplicates += 1
            continue
        prices_by_sku[sku_id] = record

    print(f"📥 Leyendo SKUs: {args.skus_csv}")
    csv_rows, csv_fieldnames = read_skus_csv(args.skus_csv, args.column)

    found, not_found, seen = [], [], set()
    csv_invalid = 0
    for row in csv_rows:
        sku_id = to_sku_id(row.get(args.column))
        if sku_id is None:
            csv_invalid += 1
            not_found.append(row)
            continue
        if sku_id in seen:
            continue
        seen.add(sku_id)
        if sku_id in prices_by_sku:
            found.append(prices_by_sku[sku_id])
        else:
            not_found.append(row)

    output_dir = os.path.dirname(args.output_prefix)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    xlsx_file = f"{args.output_prefix}_precios.xlsx"
    not_found_file = f"{args.output_prefix}_no_encontrados.csv"
    report_file = f"{args.output_prefix}_REPORT.md"

    write_prices_xlsx(xlsx_file, found)
    files = [xlsx_file]
    if not_found:
        write_csv(not_found_file, csv_fieldnames, not_found)
        files.append(not_found_file)

    stats = {
        'price_rows': len(price_records),
        'price_duplicates': price_duplicates,
        'csv_rows': len(csv_rows),
        'csv_invalid': csv_invalid,
        'found': len(found),
        'not_found': len(not_found),
    }
    generate_report(report_file, stats, files + [report_file], sheet_name)

    print(f"✅ Con precio: {len(found)} / {len(csv_rows)}")
    if not_found:
        ids = ', '.join(str(row.get(args.column, '')).strip() for row in not_found)
        print(f"⚠️  Sin precio en xlsx: {len(not_found)} ({ids})")
    print("📁 Archivos generados:")
    for path in files + [report_file]:
        print(f"   - {path}")


if __name__ == '__main__':
    main()
