#!/usr/bin/env python3
"""
xlsx_to_xls.py

Script para convertir archivos Excel .xlsx a formato Excel clásico (.xls).
Complementa el flujo de transformación de datos para integración con VTEX e-commerce platform.

Funcionalidad:
- Lee archivos .xlsx (todas las hojas por defecto, o una hoja específica con --sheet)
- Conserva los tipos de celda: números como números, texto como texto, booleanos y fechas
- Usa los valores calculados de las fórmulas (no la fórmula)
- Omite las filas vacías al final de cada hoja
- Igual que la exportación de inventario de VTEX: si una hoja supera las 65536 filas
  (límite del formato .xls), se reparte en varias hojas de 65536 filas, repitiendo la
  fila de cabeceras (primera fila) en cada una:
    * Una sola hoja de origen  -> Sheet1, Sheet2, ..., SheetN
    * Varias hojas de origen   -> NombreHoja_1, NombreHoja_2, ..., NombreHoja_N
- Las hojas que no superan el límite conservan su nombre (truncado a 31 caracteres)
- Valida los demás límites del formato .xls: 256 columnas y 32767 caracteres por celda

Dependencias:
- openpyxl: requerido para leer archivos .xlsx
- xlwt: requerido para escribir archivos .xls (clásico)

Ejecución:
    # Conversión de todas las hojas
    python3 xlsx_to_xls.py entrada.xlsx salida.xls

    # Conversión de una sola hoja
    python3 xlsx_to_xls.py entrada.xlsx salida.xls --sheet "Precios"

Ejemplo:
    python3 01_csv_to_json/xlsx_to_xls.py estoque.xlsx estoque.xls
    python3 01_csv_to_json/xlsx_to_xls.py reporte.xlsx reporte.xls --sheet Hoja1
"""
import argparse
import datetime
import decimal
import os
import sys

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import xlwt
    XLWT_AVAILABLE = True
except ImportError:
    XLWT_AVAILABLE = False

XLS_MAX_ROWS = 65536
XLS_MAX_COLUMNS = 256
XLS_MAX_CELL_LENGTH = 32767
XLS_MAX_SHEET_NAME_LENGTH = 31
# Múltiplo de 32 (tamaño de bloque de filas de xlwt)
FLUSH_EVERY_ROWS = 4096


def check_dependencies():
    """Verifica que las dependencias requeridas estén instaladas."""
    missing = []
    if not OPENPYXL_AVAILABLE:
        missing.append("openpyxl")
    if not XLWT_AVAILABLE:
        missing.append("xlwt")
    if missing:
        raise ImportError(
            f"Faltan dependencias: {', '.join(missing)}.\n"
            f"Instálalas con: pip install {' '.join(missing)}"
        )


class XlsWriter:
    """Escribe hojas en un libro .xls con nombres válidos y tipos de celda conservados."""

    def __init__(self):
        self.workbook = xlwt.Workbook(encoding="utf-8")
        self.used_names = set()
        self.datetime_style = xlwt.easyxf(num_format_str="yyyy-mm-dd hh:mm:ss")
        self.date_style = xlwt.easyxf(num_format_str="yyyy-mm-dd")
        self.time_style = xlwt.easyxf(num_format_str="hh:mm:ss")

    def _unique_name(self, sheet_name):
        """
        Genera un nombre de hoja válido para .xls (máx. 31 caracteres, único sin
        distinguir mayúsculas). Si se trunca, se conserva el sufijo final (_N).
        """
        name = str(sheet_name) or "Sheet"
        if len(name) > XLS_MAX_SHEET_NAME_LENGTH:
            head, sep, tail = name.rpartition("_")
            if sep and tail.isdigit():
                suffix = f"_{tail}"
                name = head[:XLS_MAX_SHEET_NAME_LENGTH - len(suffix)] + suffix
            else:
                name = name[:XLS_MAX_SHEET_NAME_LENGTH]
        base = name
        counter = 1
        while name.lower() in self.used_names:
            suffix = f"~{counter}"
            name = base[:XLS_MAX_SHEET_NAME_LENGTH - len(suffix)] + suffix
            counter += 1
        self.used_names.add(name.lower())
        return name

    def add_sheet(self, sheet_name):
        """Crea una hoja nueva con un nombre válido y único."""
        return self.workbook.add_sheet(self._unique_name(sheet_name))

    def rename_sheet(self, sheet, new_name):
        """Renombra una hoja ya creada (usado cuando una hoja resulta repartida)."""
        old_lower = sheet.name.lower()
        self.used_names.discard(old_lower)
        new_name = self._unique_name(new_name)
        # xlwt no expone renombrar hojas: su índice interno de nombres (usado para
        # detectar duplicados en add_sheet) se actualiza a mano. xlwt 1.3.0 es la
        # versión final de la librería, por lo que este atributo no cambiará.
        name_index = self.workbook._Workbook__worksheet_idx_from_name
        name_index[new_name.lower()] = name_index.pop(old_lower)
        sheet.set_name(new_name)

    def write_row(self, sheet, row_index, row):
        """Escribe una fila conservando el tipo de cada celda."""
        for col_index, value in enumerate(row):
            if value is None or value == "":
                continue
            if isinstance(value, bool) or isinstance(value, (int, float)):
                sheet.write(row_index, col_index, value)
            elif isinstance(value, decimal.Decimal):
                sheet.write(row_index, col_index, float(value))
            elif isinstance(value, datetime.datetime):
                sheet.write(row_index, col_index, value, self.datetime_style)
            elif isinstance(value, datetime.date):
                sheet.write(row_index, col_index, value, self.date_style)
            elif isinstance(value, datetime.time):
                sheet.write(row_index, col_index, value, self.time_style)
            else:
                sheet.write(row_index, col_index, str(value))

    def save(self, output_file):
        self.workbook.save(output_file)


def validate_row(sheet_name, row_number, row):
    """
    Valida que una fila quepa dentro de los límites del formato .xls clásico.

    :param sheet_name: Nombre de la hoja de origen
    :param row_number: Número de fila en la hoja de origen (1-based)
    :param row: Tupla de valores de la fila
    """
    last_column = max((i for i, v in enumerate(row) if v is not None and v != ""), default=-1)
    if last_column + 1 > XLS_MAX_COLUMNS:
        raise ValueError(
            f"La hoja '{sheet_name}' tiene datos en la columna {last_column + 1} (fila {row_number}); "
            f"el formato .xls clásico soporta máximo {XLS_MAX_COLUMNS} columnas."
        )
    for col_index, value in enumerate(row, start=1):
        if isinstance(value, str) and len(value) > XLS_MAX_CELL_LENGTH:
            raise ValueError(
                f"La celda fila {row_number}, columna {col_index} de la hoja '{sheet_name}' "
                f"tiene {len(value)} caracteres (máximo .xls: {XLS_MAX_CELL_LENGTH})."
            )


def is_empty_row(row):
    """Indica si una fila no tiene ningún valor."""
    return all(value is None or value == "" for value in row)


def convert_sheet(source_ws, source_name, writer, split_prefix):
    """
    Copia una hoja del .xlsx al libro .xls, repartiéndola en varias hojas de
    XLS_MAX_ROWS filas (con la cabecera repetida) si supera el límite.

    :param source_ws: Hoja openpyxl de origen (modo read_only)
    :param source_name: Nombre de la hoja de origen
    :param writer: XlsWriter de destino
    :param split_prefix: Prefijo de nombre para las hojas repartidas ('Sheet' o 'Nombre_')
    :return: Lista de tuplas (nombre_hoja_destino, filas_escritas)
    """
    target = writer.add_sheet(source_name)
    parts = [[target, 0]]
    header = None
    pending_empty = 0
    row_index = 0

    def emit(row):
        nonlocal target, row_index
        if row_index == XLS_MAX_ROWS:
            if len(parts) == 1:
                writer.rename_sheet(parts[0][0], f"{split_prefix}1")
            target = writer.add_sheet(f"{split_prefix}{len(parts) + 1}")
            parts.append([target, 0])
            writer.write_row(target, 0, header)
            row_index = 1
        writer.write_row(target, row_index, row)
        row_index += 1
        parts[-1][1] = row_index
        if row_index % FLUSH_EVERY_ROWS == 0:
            # Serializa las filas ya escritas a un archivo temporal para no
            # mantener todo el libro en memoria (archivos de cientos de miles de filas)
            target.flush_row_data()

    for row_number, row in enumerate(source_ws.iter_rows(values_only=True), start=1):
        if is_empty_row(row):
            pending_empty += 1
            continue
        validate_row(source_name, row_number, row)
        if header is None:
            header = row
            # Filas vacías antes de la cabecera se conservan tal cual
            row_index = pending_empty
            pending_empty = 0
            emit(row)
            continue
        for _ in range(pending_empty):
            emit(())
        pending_empty = 0
        emit(row)

    return [(sheet.name, rows) for sheet, rows in parts]


def convert_xlsx_to_xls(input_file, output_file, sheet=None):
    """
    Convierte un archivo .xlsx a .xls clásico.

    :param input_file: Ruta del .xlsx de entrada
    :param output_file: Ruta del .xls de salida
    :param sheet: Nombre de la hoja a convertir. Si es None, se convierten todas
    :return: dict {hoja_origen: [(hoja_destino, filas), ...]}
    """
    workbook = openpyxl.load_workbook(input_file, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise ValueError(
                    f"La hoja '{sheet}' no existe. Hojas disponibles: {', '.join(workbook.sheetnames)}"
                )
            sheet_names = [sheet]
        else:
            sheet_names = workbook.sheetnames

        writer = XlsWriter()
        summary = {}
        single_source = len(sheet_names) == 1
        for name in sheet_names:
            split_prefix = "Sheet" if single_source else f"{name}_"
            summary[name] = convert_sheet(workbook[name], name, writer, split_prefix)
    finally:
        workbook.close()

    writer.save(output_file)
    return summary


def main():
    """Función principal que maneja la conversión de .xlsx a .xls."""
    parser = argparse.ArgumentParser(
        description="Convierte archivos Excel .xlsx a formato Excel clásico (.xls).\n"
                    "Las hojas de más de 65536 filas se reparten en varias hojas con la\n"
                    "cabecera repetida, igual que la exportación de inventario de VTEX.",
        formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument(
        "input_file",
        help="Archivo .xlsx de entrada"
    )
    parser.add_argument(
        "output_file",
        help="Archivo .xls de salida"
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Nombre de la hoja a convertir. Si se omite, se convierten todas las hojas"
    )

    args = parser.parse_args()

    try:
        if not os.path.exists(args.input_file):
            raise FileNotFoundError(args.input_file)

        check_dependencies()

        summary = convert_xlsx_to_xls(args.input_file, args.output_file, sheet=args.sheet or None)

        print(f"Successfully converted {args.input_file} to {args.output_file} (xls)")
        total_sheets = 0
        for source_name, parts in summary.items():
            if len(parts) > 1:
                print(f"  Hoja '{source_name}' repartida en {len(parts)} hojas (cabecera repetida):")
            for target_name, rows in parts:
                indent = "    " if len(parts) > 1 else "  "
                print(f"{indent}Hoja '{target_name}': {rows} filas")
            total_sheets += len(parts)
        print(f"Total sheets: {total_sheets}")

    except ImportError as e:
        sys.stderr.write(f"Error de dependencias: {e}\n")
        sys.exit(1)
    except FileNotFoundError:
        sys.stderr.write(f"Error: Archivo no encontrado - {args.input_file}\n")
        sys.exit(1)
    except ValueError as e:
        sys.stderr.write(f"Error: {e}\n")
        sys.exit(1)
    except Exception as e:
        sys.stderr.write(f"Error al convertir archivo: {e}\n")
        sys.exit(1)


if __name__ == "__main__":
    main()
